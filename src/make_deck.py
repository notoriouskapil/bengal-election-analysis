"""Build the deck twice from one content source.

    outputs/Bengal_2021_vs_2026.pptx           16:9, for sending to people
    outputs/Bengal_2021_vs_2026_carousel.pdf   1080x1080, for LinkedIn

Both read SLIDES below, so the two files cannot drift apart.
Run src/make_charts.py first.

    python src/make_deck.py
"""

import textwrap
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.backends.backend_pdf import PdfPages
from matplotlib.patches import FancyBboxPatch
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
CHARTS = ROOT / "outputs" / "charts"
OUT = ROOT / "outputs"

BJP, TMC, BLUE = "#eb6834", "#0f6b42", "#2a78d6"
INK, SUB, GREY = "#16150f", "#52514e", "#8a8782"
RULE, SURFACE, PANEL = "#e8e7e3", "#fcfcfb", "#f4f3f0"
FONT = ["Helvetica Neue", "Helvetica", "Arial", "DejaVu Sans"]
MONO = ["Menlo", "DejaVu Sans Mono", "monospace"]


def rgb(h):
    return RGBColor(int(h[1:3], 16), int(h[3:5], 16), int(h[5:7], 16))


SLIDES = [
    dict(kind="title",
         title="West Bengal\n2021 vs 2026",
         sub="294 constituencies  ·  5,052 candidacies\ntwo elections, compared seat by seat",
         foot="Kapil   ·   github.com/notoriouskapil/bengal-election-analysis",
         notes="Two West Bengal assembly elections, built from the Election Commission's own "
               "statistical reports. I built the dataset, wrote the queries, made the charts."),
    dict(kind="bullets",
         title="Four numbers",
         bullets=[("BJP: 46.2% of votes", "70.6% of seats", BJP),
                  ("129 seats TMC → BJP", "Zero returned", BJP),
                  ("Turnout 82.2% → 93.6%", "5m names left the roll", BLUE),
                  ("TMC fell in all 8 regions", "worst: 50% → 36%", TMC)],
         notes="Findings one and two are first-past-the-post doing what it does. "
               "Finding three is the one that changes how you read the headline turnout figure."),
    dict(kind="data",
         title="Where this came from",
         body="15 ECI statistical workbooks covering both elections, reshaped into seven "
              "tables and checked back against the totals ECI publishes in its own "
              "Highlight report.",
         checks=[("electors, 293 polled seats", "68,125,496"),
                 ("valid votes", "63,258,138"),
                 ("NOTA votes", "494,932"),
                 ("contestants", "2,920")],
         foot="Every figure reconciles to the published report. Nine defects in the source "
              "data found and corrected — including constituency names that match 0 of "
              "294 across the two years.",
         notes="The assertions matter more than the cleaning. If the ECI reissues a workbook "
               "the pipeline fails loudly instead of quietly producing different numbers."),
    dict(kind="finding", tag="FINDING 1",
         title="Most of the turnout rise is a smaller denominator, not more voters",
         chart="06_turnout_vs_roll.png",
         body="Run 2026's votes against the 2021 roll and turnout reads 87.1%, not "
              "93.6%. More than half the rise is 5 million names leaving the roll. "
              "Cause unestablished.",
         notes="82.15% in 2021. 93.58% reported in 2026. 87.12% if you hold the roll fixed. "
               "The roll fell by 5.05 million across 293 seats and 241 of them lost voters. "
               "What drove the revision is not something this data can answer."),
    dict(kind="finding", tag="FINDING 2",
         title="46.2% of the vote. 70.6% of the seats.",
         chart="02_vote_vs_seat_share.png",
         body="Under first-past-the-post a plurality nearly everywhere becomes a supermajority. "
              "BJP went from 38.4% of the vote to 46.2%, and gained 130 seats.",
         notes="Seat share is over 293, not 294: Falta held no poll in 2026. "
               "TMC sits on the other side of the same mechanism — 41.1% of votes, 27.3% of seats."),
    dict(kind="finding", tag="FINDING 3",
         title="129 seats went one way. None came back.",
         chart="03_flip_matrix.png",
         body="136 of 293 seats changed hands and not one moved to TMC. BJP took 129 from TMC "
              "and one from an independent; the other six went to smaller parties.",
         notes="Not a single seat flipped against the tide. The two-way flip matrix is the "
               "only join that answers this — a name join silently fans out."),
    dict(kind="closing",
         title="What this can't tell you",
         cols=[("LIMITATIONS", ["District mapping derived, not sourced",
                                "No cross-year candidate tracking — 543 of 2,079 names match",
                                "Nothing here is causal"]),
               ("NEXT", ["Find what drove the roll revision",
                         "Add 2016 for a three-election trend"])],
         notes="The roll revision is the open question. I can measure the effect. "
               "I cannot say what caused it, and I am not going to guess on a slide."),
    dict(kind="appendix", chart="01_seat_reversal.png",
         title="TMC 215 → 80. BJP 77 → 207.",
         body="TMC's vote share fell from 48.5% to 41.1% and it lost 135 seats. Seats "
              "won, with each year's vote share beneath.",
         notes="The slope chart is the setup for the vote-to-seat slide, not a separate finding."),
    dict(kind="appendix", chart="04_closest_seats_2026.png",
         title="BJP won 6 of the 8 seats decided by less than 1%",
         body="Ranked on margin as a percentage of votes polled — the measure that stays "
              "comparable across seats of different size.",
         notes="Rajarhat New Town came down to 316 votes."),
    dict(kind="appendix", chart="05_regional_swing.png",
         title="TMC's vote share fell in every region",
         body="From 52% down to 48% in South 24 Parganas, and 50% down to 36% in "
              "Murshidabad-Nadia. Regions are derived from district, not sourced.",
         notes="Eight regions built from the district mapping, which is inferred. "
               "A boundary error would move a region's number but not its sign."),
]


# ------------------------------------------------------------ schema figure

def draw_schema(path):
    fig = plt.figure(figsize=(9.0, 5.1), dpi=200, facecolor=SURFACE)
    ax = fig.add_axes([0, 0, 1, 1]); ax.set_xlim(0, 100); ax.set_ylim(-10, 52)
    ax.axis("off")

    def box(x, y, w, h, label, rows, colour, bold=False):
        ax.add_patch(FancyBboxPatch((x, y), w, h, boxstyle="round,pad=0.4,rounding_size=1.1",
                                    linewidth=1.6 if bold else 1.1,
                                    edgecolor=colour, facecolor=PANEL if bold else SURFACE))
        ax.text(x + w / 2, y + h * 0.60, label, ha="center", va="center",
                fontsize=10.5 if bold else 9.5, color=INK,
                fontweight="bold" if bold else "normal", family=MONO)
        ax.text(x + w / 2, y + h * 0.24, rows, ha="center", va="center",
                fontsize=8.5, color=SUB)

    box(36, 21, 28, 10, "constituencies", "294 rows  ·  PK ac_no", BLUE, bold=True)
    leaves = [(2, 38, "results", "5,052"), (35.5, 38, "winners", "587"), (69, 38, "electorate", "588"),
              (2, 4, "nota", "588"), (35.5, 4, "seat_status", "588"), (69, 4, "name_flags", "8")]
    for x, y, name, n in leaves:
        colour = TMC if y > 20 else BJP
        box(x, y, 29, 9, name, f"{n} rows  ·  FK ac_no", colour)
        cx, cy = x + 14.5, y + (0 if y > 20 else 9)
        ax.plot([cx, 50], [cy, 31 if y > 20 else 21], color=RULE, linewidth=1.3, zorder=0)

    ax.text(50, -3.4, "v_results   ·   v_winners", ha="center", fontsize=9,
            color=SUB, family=MONO)
    ax.text(50, -6.8, "views with district, region and reservation pre-joined",
            ha="center", fontsize=8, color=GREY)
    fig.savefig(path, dpi=200, facecolor=SURFACE)
    plt.close(fig)
    print(f"  wrote {path.relative_to(ROOT)}")


# ------------------------------------------------------------ pptx

W, H = Inches(13.333), Inches(7.5)


def tb(slide, x, y, w, h, text, size, colour=INK, bold=False, font=None,
       align=PP_ALIGN.LEFT, space=0.0, line=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, chunk in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if space:
            p.space_after = Pt(space)
        if line:
            p.line_spacing = line
        r = p.add_run(); r.text = chunk
        r.font.size = Pt(size); r.font.bold = bold
        r.font.color.rgb = rgb(colour)
        r.font.name = (font or FONT[0])
    return box


def band(slide, x, y, w, h, colour):
    from pptx.enum.shapes import MSO_SHAPE
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb = rgb(colour)
    s.line.fill.background(); s.shadow.inherit = False
    return s


def picture(slide, img, x, y, max_w, max_h):
    from PIL import Image
    iw, ih = Image.open(img).size
    scale = min(max_w / (iw / 200), max_h / (ih / 200))
    w, h = (iw / 200) * scale, (ih / 200) * scale
    slide.shapes.add_picture(str(img), Inches(x + (max_w - w) / 2),
                             Inches(y + (max_h - h) / 2), Inches(w), Inches(h))


def build_pptx(path):
    prs = Presentation(); prs.slide_width, prs.slide_height = W, H
    blank = prs.slide_layouts[6]

    for s in SLIDES:
        sl = prs.slides.add_slide(blank)
        band(sl, 0, 0, 13.333, 7.5, SURFACE)
        k = s["kind"]

        if k == "title":
            band(sl, 0.9, 2.05, 1.5, 0.07, BJP)
            tb(sl, 0.9, 2.45, 11.5, 2.2, s["title"], 46, INK, bold=True, line=1.05)
            tb(sl, 0.9, 4.75, 11.5, 0.8, s["sub"], 15, SUB)
            band(sl, 0.9, 6.25, 11.5, 0.012, RULE)
            tb(sl, 0.9, 6.5, 11.5, 0.5, s["foot"], 11.5, GREY)

        elif k == "bullets":
            tb(sl, 0.9, 0.72, 11.5, 0.8, s["title"], 30, INK, bold=True)
            band(sl, 0.9, 1.62, 1.5, 0.05, BJP)
            y = 2.2
            for head, tail, colour in s["bullets"]:
                band(sl, 0.9, y + 0.08, 0.055, 0.62, colour)
                tb(sl, 1.25, y, 5.4, 0.7, head, 19, INK)
                tb(sl, 6.6, y, 6.0, 0.7, tail, 19, colour, bold=True)
                y += 1.08

        elif k == "data":
            tb(sl, 0.9, 0.72, 11.5, 0.8, s["title"], 30, INK, bold=True)
            band(sl, 0.9, 1.62, 1.5, 0.05, BJP)
            tb(sl, 0.9, 2.05, 6.2, 2.2, s["body"], 16, SUB, line=1.45)
            band(sl, 7.5, 1.95, 4.95, 3.1, PANEL)
            y = 2.25
            for label, value in s["checks"]:
                tb(sl, 7.85, y, 0.5, 0.4, "ok", 12, TMC, bold=True, font=MONO[0])
                tb(sl, 8.45, y, 2.6, 0.4, label, 12, SUB, font=MONO[0])
                tb(sl, 10.3, y, 1.95, 0.4, value, 12, INK, font=MONO[0], align=PP_ALIGN.RIGHT)
                y += 0.62
            tb(sl, 0.9, 5.55, 11.5, 0.8, s["foot"], 12.5, GREY)

        elif k == "schema":
            tb(sl, 0.9, 0.62, 11.5, 0.7, s["title"], 30, INK, bold=True)
            tb(sl, 0.9, 1.42, 9.6, 0.8, s["body"], 14.5, SUB, line=1.4)
            picture(sl, CHARTS / "00_schema.png", 0.9, 2.35, 11.5, 4.5)

        elif k in ("finding", "appendix"):
            if s.get("tag"):
                tb(sl, 0.9, 0.55, 6.0, 0.35, s["tag"], 11, BJP, bold=True)
                tb(sl, 0.9, 0.95, 11.6, 1.0, s["title"], 25, INK, bold=True, line=1.15)
            else:
                tb(sl, 0.9, 0.62, 11.6, 1.0, s["title"], 25, INK, bold=True, line=1.15)
            picture(sl, CHARTS / s["chart"], 0.7, 1.95, 8.35, 4.8)
            band(sl, 9.35, 2.05, 3.1, 0.045, RULE)
            tb(sl, 9.35, 2.3, 3.1, 3.6, s["body"], 13.5, SUB, line=1.45)

        elif k == "closing":
            tb(sl, 0.9, 0.72, 11.5, 0.8, s["title"], 28, INK, bold=True)
            band(sl, 0.9, 1.62, 1.5, 0.05, BJP)
            for i, (head, items) in enumerate(s["cols"]):
                x = 0.9 + i * 5.9
                tb(sl, x, 2.15, 5.3, 0.35, head, 11.5, BJP, bold=True)
                y = 2.65
                for it in items:
                    tb(sl, x, y, 5.3, 1.1, "·  " + it, 14.5, SUB, line=1.4)
                    y += 0.40 + 0.32 * (len(it) // 50)

        if s.get("notes"):
            sl.notes_slide.notes_text_frame.text = s["notes"]

    prs.save(path)
    print(f"  wrote {path.relative_to(ROOT)}")


# ------------------------------------------------------------ square PDF

SQ = 10.8  # inches at 100 dpi -> 1080 x 1080


def pdf_text(fig, x, y, text, size, colour=INK, weight="normal", family=None,
             ha="left", va="top", wrap=None, leading=1.45):
    lines = []
    for para in text.split("\n"):
        lines.extend(textwrap.wrap(para, wrap) if wrap else [para])
    step = size * leading / 72 / SQ
    for i, ln in enumerate(lines):
        fig.text(x, y - i * step, ln, fontsize=size, color=colour, fontweight=weight,
                 family=family or FONT, ha=ha, va=va)
    return y - len(lines) * step


def pdf_rule(fig, x, y, w, colour=RULE, h=0.0035):
    fig.patches.append(plt.Rectangle((x, y), w, h, transform=fig.transFigure,
                                     facecolor=colour, edgecolor="none"))


def pdf_chart(fig, img, x, y, w, h):
    im = plt.imread(img)
    ih, iw = im.shape[0], im.shape[1]
    scale = min(w / iw, h / ih)
    cw, ch = iw * scale, ih * scale
    ax = fig.add_axes([x + (w - cw) / 2, y + (h - ch) / 2, cw, ch])
    ax.imshow(im); ax.axis("off")


def build_pdf(path):
    with PdfPages(path) as pdf:
        for s in SLIDES:
            fig = plt.figure(figsize=(SQ, SQ), dpi=100, facecolor=SURFACE)
            k = s["kind"]
            M = 0.075

            if k == "title":
                pdf_rule(fig, M, 0.70, 0.11, BJP, h=0.007)
                y = pdf_text(fig, M, 0.665, s["title"], 52, INK, "bold", leading=1.18)
                pdf_text(fig, M, y - 0.045, s["sub"], 16.5, SUB, leading=1.65)
                pdf_rule(fig, M, 0.115, 1 - 2 * M)
                pdf_text(fig, M, 0.09, s["foot"], 12, GREY, leading=1.7)

            elif k == "bullets":
                pdf_text(fig, M, 0.91, s["title"], 34, INK, "bold")
                pdf_rule(fig, M, 0.855, 0.11, BJP, h=0.006)
                y = 0.76
                for head, tail, colour in s["bullets"]:
                    fig.patches.append(plt.Rectangle((M, y - 0.085), 0.006, 0.085,
                                       transform=fig.transFigure, facecolor=colour,
                                       edgecolor="none"))
                    pdf_text(fig, M + 0.028, y, head, 21, INK)
                    pdf_text(fig, M + 0.028, y - 0.043, tail, 24, colour, "bold")
                    y -= 0.165

            elif k == "data":
                pdf_text(fig, M, 0.91, s["title"], 34, INK, "bold")
                pdf_rule(fig, M, 0.855, 0.11, BJP, h=0.006)
                pdf_text(fig, M, 0.79, s["body"], 17, SUB, wrap=50, leading=1.62)
                fig.patches.append(plt.Rectangle((M, 0.30), 1 - 2 * M, 0.30,
                                   transform=fig.transFigure, facecolor=PANEL, edgecolor="none"))
                y = 0.555
                for label, value in s["checks"]:
                    pdf_text(fig, M + 0.03, y, "ok", 14, TMC, "bold", family=MONO)
                    pdf_text(fig, M + 0.075, y, label, 14, SUB, family=MONO)
                    pdf_text(fig, 1 - M - 0.03, y, value, 14, INK, family=MONO, ha="right")
                    y -= 0.062
                pdf_text(fig, M, 0.24, s["foot"], 13, GREY, wrap=62, leading=1.55)

            elif k == "schema":
                pdf_text(fig, M, 0.93, s["title"], 34, INK, "bold")
                pdf_text(fig, M, 0.865, s["body"], 16, SUB, wrap=56, leading=1.6)
                pdf_chart(fig, CHARTS / "00_schema.png", M, 0.20, 1 - 2 * M, 0.60)

            elif k in ("finding", "appendix"):
                if s.get("tag"):
                    pdf_text(fig, M, 0.955, s["tag"], 12.5, BJP, "bold")
                    y = pdf_text(fig, M, 0.925, s["title"], 29, INK, "bold",
                                 wrap=40, leading=1.26)
                else:
                    y = pdf_text(fig, M, 0.945, s["title"], 29, INK, "bold",
                                 wrap=40, leading=1.26)
                top = 0.255
                pdf_chart(fig, CHARTS / s["chart"], M, top, 1 - 2 * M, y - top - 0.018)
                pdf_rule(fig, M, top - 0.05, 1 - 2 * M)
                pdf_text(fig, M, top - 0.085, s["body"], 16.5, SUB, wrap=60, leading=1.58)

            elif k == "closing":
                pdf_text(fig, M, 0.92, s["title"], 30, INK, "bold", wrap=34, leading=1.24)
                pdf_rule(fig, M, 0.79, 0.11, BJP, h=0.006)
                y = 0.70
                for head, items in s["cols"]:
                    pdf_text(fig, M, y, head, 13.5, BJP, "bold")
                    y -= 0.05
                    for it in items:
                        y = pdf_text(fig, M, y, "·  " + it, 17, SUB, wrap=52,
                                     leading=1.55) - 0.022
                    y -= 0.045

            pdf.savefig(fig, facecolor=SURFACE)
            plt.close(fig)
    print(f"  wrote {path.relative_to(ROOT)}")


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    draw_schema(CHARTS / "00_schema.png")
    build_pptx(OUT / "Bengal_2021_vs_2026.pptx")
    build_pdf(OUT / "Bengal_2021_vs_2026_carousel.pdf")
    print(f"done, {len(SLIDES)} slides")


if __name__ == "__main__":
    main()
